# -*- coding: utf-8 -*-
"""
Сервис семантической сверки текстовой информации.
Бэкенд на FastAPI.
"""

import io
import math
import os
import re
from pathlib import Path

from fastapi import FastAPI, UploadFile, File
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

import docx
import numpy as np

# ======================================================================
#  Параметры математической модели (раздел 2.3 ВКР)
#  S = ALPHA*s_лекс + BETA*s_сем + GAMMA*s_стил
# ======================================================================
ALPHA = 0.15
BETA  = 0.55
GAMMA = 0.30
DEFAULT_TAU = 0.75
MIN_PARAGRAPH_LEN = 40

RUSSIAN_STOPWORDS = set("""
и в во не что он на я с со как а то все она так его но да ты к у же вы за бы
по только ее мне было вот от меня еще нет о из ему теперь когда даже ну вдруг
ли если уже или ни быть был него до вас нибудь опять уж вам ведь там потом себя
ничего ей может они тут где есть надо ней для мы тебя их чем была сам чтоб без
будто чего раз тоже себе под будет ж тогда кто этот того потому этого какой
совсем ним здесь этом один почти мой тем чтобы нее сейчас были куда зачем всех
никогда можно при наконец два об другой хоть после над больше тот через эти нас
про всего них какая много разве три эту моя впрочем хорошо свою этой перед иногда
лучше чуть том нельзя такой им более всегда конечно всю между это
""".split())

try:
    import pymorphy2
    _morph = pymorphy2.MorphAnalyzer()
except Exception:
    _morph = None

_model = None
def get_model():
    global _model
    if _model is None:
        from sentence_transformers import SentenceTransformer
        _model = SentenceTransformer("paraphrase-multilingual-MiniLM-L12-v2")
    return _model


def extract_sections(file_bytes: bytes):
    doc = docx.Document(io.BytesIO(file_bytes))
    sections = []
    cur_title = "Документ"
    cur_paras = []
    started = False

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        is_heading = style.startswith("heading") or style.startswith("заголовок")

        if is_heading:
            if started and cur_paras:
                sections.append({"title": cur_title, "paragraphs": cur_paras})
            cur_title = text
            cur_paras = []
            started = True
        else:
            if len(text) >= MIN_PARAGRAPH_LEN:
                cur_paras.append(text)

    if started and cur_paras:
        sections.append({"title": cur_title, "paragraphs": cur_paras})

    if not sections:
        all_paras = [p.text.strip() for p in doc.paragraphs
                     if len(p.text.strip()) >= MIN_PARAGRAPH_LEN]
        if all_paras:
            sections = [{"title": "Документ", "paragraphs": all_paras}]

    return sections


def preprocess(text: str):
    text = text.lower()
    text = re.sub(r"[^а-яёa-z0-9\s]", " ", text)
    tokens = [t for t in text.split() if t and t not in RUSSIAN_STOPWORDS]
    if _morph:
        tokens = [_morph.parse(t)[0].normal_form for t in tokens]
    return tokens


def levenshtein(a, b):
    m, n = len(a), len(b)
    if m == 0: return n
    if n == 0: return m
    prev = list(range(n + 1))
    for i in range(1, m + 1):
        cur = [i] + [0] * n
        ai = a[i - 1]
        for j in range(1, n + 1):
            cost = 0 if ai == b[j - 1] else 1
            cur[j] = min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + cost)
        prev = cur
    return prev[n]


def s_lex(tok1, tok2):
    if not tok1 and not tok2:
        return 1.0
    d = levenshtein(tok1, tok2)
    return 1.0 - d / max(len(tok1), len(tok2), 1)


def stylo_vector(text: str):
    words = text.split()
    if not words:
        return [0.0] * 5
    avg_word = sum(len(w) for w in words) / len(words)
    n_sent = max(text.count(".") + text.count("!") + text.count("?"), 1)
    avg_sent = len(words) / n_sent
    lex_div = len(set(words)) / len(words)
    punct = sum(1 for c in text if c in ",.;:!?-—()") / max(len(text), 1)
    stop = sum(1 for w in words if w.lower() in RUSSIAN_STOPWORDS) / len(words)
    return [avg_word, avg_sent, lex_div, punct, stop]


def gaussian_kernel(v1, v2, h):
    dist_sq = sum((a - b) ** 2 for a, b in zip(v1, v2))
    return math.exp(-dist_sq / (2 * h * h))


# ======================================================================
#  FastAPI
# ======================================================================
app = FastAPI(title="Сервис семантической сверки текстов")


@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/api/parse")
async def api_parse(main_file: UploadFile = File(...),
                    compare_file: UploadFile = File(...)):
    main_bytes = await main_file.read()
    comp_bytes = await compare_file.read()
    return {
        "main":    {"filename": main_file.filename,    "sections": extract_sections(main_bytes)},
        "compare": {"filename": compare_file.filename, "sections": extract_sections(comp_bytes)},
    }


class CompareRequest(BaseModel):
    main_sections: list
    compare_sections: list
    threshold: float = DEFAULT_TAU


@app.post("/api/compare")
def api_compare(req: CompareRequest):
    paras1 = [(s["title"], p) for s in req.main_sections    for p in s["paragraphs"]]
    paras2 = [(s["title"], p) for s in req.compare_sections for p in s["paragraphs"]]

    if not paras1 or not paras2:
        return {"results": [], "pairs_compared": 0, "threshold": req.threshold}

    texts1 = [p for _, p in paras1]
    texts2 = [p for _, p in paras2]

    model = get_model()
    emb1 = model.encode(texts1, convert_to_numpy=True)
    emb2 = model.encode(texts2, convert_to_numpy=True)
    emb1 = emb1 / (np.linalg.norm(emb1, axis=1, keepdims=True) + 1e-9)
    emb2 = emb2 / (np.linalg.norm(emb2, axis=1, keepdims=True) + 1e-9)

    tok1 = [preprocess(t) for t in texts1]
    tok2 = [preprocess(t) for t in texts2]

    sv1 = [stylo_vector(t) for t in texts1]
    sv2 = [stylo_vector(t) for t in texts2]
    allv = sv1 + sv2
    cols = list(zip(*allv))
    means = [sum(c) / len(c) for c in cols]
    stds = [(sum((x - m) ** 2 for x in c) / len(c)) ** 0.5 or 1.0
            for c, m in zip(cols, means)]

    def znorm(v):
        return [(x - m) / s for x, m, s in zip(v, means, stds)]

    zv1 = [znorm(v) for v in sv1]
    zv2 = [znorm(v) for v in sv2]
    h = math.sqrt(len(means))

    results = []
    for i, (t1, p1) in enumerate(paras1):
        for j, (t2, p2) in enumerate(paras2):
            sem = max(0.0, float(np.dot(emb1[i], emb2[j])))
            lex = s_lex(tok1[i], tok2[j])
            stil = gaussian_kernel(zv1[i], zv2[j], h)
            S = ALPHA * lex + BETA * sem + GAMMA * stil
            if S >= req.threshold:
                results.append({
                    "section_1": t1, "section_2": t2,
                    "fragment_1": p1, "fragment_2": p2,
                    "score": round(S, 2),
                    "s_lex": round(lex, 2),
                    "s_sem": round(sem, 2),
                    "s_stil": round(stil, 2),
                })

    results.sort(key=lambda x: -x["score"])
    return {
        "results": results,
        "pairs_compared": len(paras1) * len(paras2),
        "threshold": req.threshold,
    }


# ----------------------------------------------------------------------
#  Отдача фронтенда — путь вычисляется относительно этого файла
# ----------------------------------------------------------------------
FRONTEND_DIR = Path(__file__).resolve().parent.parent / "frontend"
app.mount("/", StaticFiles(directory=str(FRONTEND_DIR), html=True), name="frontend")

if __name__ == "__main__":
    import os
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run(app, host="0.0.0.0", port=port)
